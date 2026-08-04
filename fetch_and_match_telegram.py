import os, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

import json
import string
import unicodedata
from pathlib import Path
from datetime import datetime

from dotenv import load_dotenv

load_dotenv()

# ========== PYROGRAM SETUP ==========

API_ID = int(os.environ["TELEGRAM_API_ID"])
API_HASH = os.environ["TELEGRAM_API_HASH"]
PHONE = os.environ["TELEGRAM_PHONE"]
GROUP_ID = int(os.environ.get("INVENTARIO_GROUP_ID", "-1004286637959"))

try:
    from pyrogram import Client
    import pyrogram.utils
    pyrogram.utils.MIN_CHANNEL_ID = -1009999999999
except ImportError:
    print("ERROR: Pyrogram no est\u00e1 instalado.  pip install pyrogram")
    sys.exit(1)

from supabase import create_client

# ========== SUPABASE CLIENTS ==========

INV_SUPABASE_URL = os.environ["INVENTARIO_SUPABASE_URL"]
INV_SUPABASE_KEY = os.environ["INVENTARIO_SUPABASE_SERVICE_KEY"]
SUPABASE_URL = os.environ["SUPABASE_URL"]
SUPABASE_KEY = os.environ["SUPABASE_SERVICE_KEY"]
AYUDAS_SUPABASE_URL = os.environ.get("AYUDAS_SUPABASE_URL") or SUPABASE_URL
AYUDAS_SUPABASE_KEY = os.environ.get("AYUDAS_SUPABASE_SERVICE_KEY") or SUPABASE_KEY

sb_inv = create_client(INV_SUPABASE_URL, INV_SUPABASE_KEY)
sb_bot = create_client(SUPABASE_URL, SUPABASE_KEY)
sb_ayudas = create_client(AYUDAS_SUPABASE_URL, AYUDAS_SUPABASE_KEY)

SEARCH_CHARS = list(string.ascii_lowercase + string.digits)

BASE_DIR = Path(__file__).parent

# ========== HELPERS ==========

def normalize_str(s):
    if not s:
        return ''
    return unicodedata.normalize('NFKD', str(s)).encode('ASCII', 'ignore').decode().strip()

def norm_username(u):
    if not u:
        return ""
    return u.lower().strip().lstrip("@")

# ========== TELEGRAM FETCH ==========

def fetch_all_members(app, chat_id):
    seen = {}

    chat = app.get_chat(chat_id)
    print(f"  Miembros segun Telegram: {chat.members_count}")

    print(f"  Obteniendo miembros recientes...")
    for m in app.get_chat_members(chat_id, query=""):
        if not m.user.is_bot:
            seen[m.user.id] = m.user
    print(f"    -> {len(seen)}")

    total_chars = len(SEARCH_CHARS)
    for idx, ch in enumerate(SEARCH_CHARS, 1):
        print(f"  Buscando '{ch}' ({idx}/{total_chars})...", end=" ")
        try:
            for m in app.get_chat_members(chat_id, query=ch):
                if not m.user.is_bot:
                    seen[m.user.id] = m.user
        except Exception:
            pass
        print(f"{len(seen)}")

    return list(seen.values()), chat.members_count

# ========== LOAD DB SOURCES ==========

print("Cargando bases de datos...\n")

print("  inventario_adquisiciones...")
r = sb_inv.table("inventario_adquisiciones").select("telegram_username,nombre,dni,telegram_user_id").execute()
inv_by_username = {}
inv_by_user_id = {}
for row in r.data or []:
    u = norm_username(row.get("telegram_username"))
    if u:
        inv_by_username[u] = row
    uid = row.get("telegram_user_id")
    if uid and isinstance(uid, int) and uid > 0:
        inv_by_user_id[uid] = row
print(f"    {len(inv_by_username)} por username, {len(inv_by_user_id)} por user_id")

print("  reparticion_vaquita...")
r = sb_inv.table("reparticion_vaquita").select("telegram_username,nombres,documento").execute()
rep_by_username = {}
for row in r.data or []:
    u = norm_username(row.get("telegram_username"))
    if u:
        rep_by_username[u] = row
print(f"    {len(rep_by_username)} por username")

print("  usuarios_funlidi...")
r = sb_bot.table("usuarios_funlidi").select("telegram_username,nombres_completos,numero_documento,telegram_user_id").execute()
bot_by_username = {}
bot_by_user_id = {}
for row in r.data or []:
    u = norm_username(row.get("telegram_username"))
    if u:
        bot_by_username[u] = row
    uid = row.get("telegram_user_id")
    if uid and isinstance(uid, int) and uid > 0:
        bot_by_user_id[uid] = row
print(f"    {len(bot_by_username)} por username")

print("  ayudas_humanitarias...")
r = sb_ayudas.table("ayudas_humanitarias").select("telegram_username,nombre,dni").execute()
ayu_by_username = {}
for row in r.data or []:
    u = norm_username(row.get("telegram_username"))
    if u:
        ayu_by_username[u] = row
print(f"    {len(ayu_by_username)} por username")

print("  Farley CRM...")
farley_by_telegram = {}
farley_path = BASE_DIR / "data_to_show" / "dashboard-data.json"
if farley_path.exists():
    with open(farley_path, "r", encoding="utf-8") as f:
        crm_data = json.load(f)
    for member in crm_data.get("members", []):
        tel = member.get("telegram", "")
        if tel:
            tel_clean = tel.lstrip("@").lower()
            if tel_clean not in farley_by_telegram:
                farley_by_telegram[tel_clean] = member
print(f"    {len(farley_by_telegram)} por telegram")

print("  newlist.txt (fallback)...")
newlist_by_telegram = {}
newlist_path = r'C:\Users\amazi\Desktop\mariaelvira\newlist.txt'
if os.path.exists(newlist_path):
    with open(newlist_path, "r", encoding="utf-8") as f:
        for line in f:
            m = __import__('re').match(r'@(\w+)', line)
            if m:
                pass  # We'll parse blocks properly instead
    # Simple block-based parser
    with open(newlist_path, "r", encoding="utf-8") as f:
        text = f.read()
    blocks = __import__('re').split(r'\n\s*-\s*\n', text)
    for block in blocks:
        lines = [l.strip() for l in block.strip().split('\n') if l.strip()]
        name = ''
        doc_id = ''
        tel = ''
        for i, line in enumerate(lines):
            s = line.strip()
            if s.startswith('@') and not s.startswith('@ '):
                tel = s.lstrip('@').lower()
                if not name and i > 0:
                    candidate = lines[i - 1].strip()
                    if not candidate.startswith('@') and ':' not in candidate and not candidate.startswith('+'):
                        name = candidate
            elif __import__('re').search(r'c[eé]dula', s, __import__('re').IGNORECASE):
                m2 = __import__('re').search(r'(\d[\d\s]*\d)', s)
                if m2:
                    doc_id = m2.group(1).strip()
        if tel and tel not in newlist_by_telegram:
            newlist_by_telegram[tel] = {'nombre': name, 'documento': doc_id}
    # Fallback overrides for known entries
    fallback = {
        'chelym1968':    ('Maria A. Aviles Martinez',     '2548364'),
        'alexabel2087':  ('Alexis Dilone Guzman',        '7321849'),
        'isab3l2086':    ('Isabel Sanchez Sierra',       '2016042'),
        'emunah2534':    ('Elvira Cuevas Romero',        '1381327'),
        'magdapaz52':    ('Magda Irizarry Ceballos',     '922987'),
        'lunamar26':     ('Maria Elena Celestino Iglesias', 'O9326280X'),
        'natividad1225': ('Nayda T. Tefel Isern',       '1330220'),
        'lisette_torres':('Lisette Maria Torres Santiago','1121242'),
    }
    for u, (n, d) in fallback.items():
        newlist_by_telegram[u] = {'nombre': n, 'documento': d}
print(f"    {len(newlist_by_telegram)} entradas")

# ========== MATCH FUNCTION ==========

def lookup_member(user):
    uname = norm_username(user.username)

    # Priority 1: inventario_adquisiciones
    if uname and uname in inv_by_username:
        row = inv_by_username[uname]
        return {"name": normalize_str(row.get("nombre", "")), "doc": normalize_str(row.get("dni", "")), "src": "INVENTARIO"}
    if user.id > 0 and user.id in inv_by_user_id:
        row = inv_by_user_id[user.id]
        return {"name": normalize_str(row.get("nombre", "")), "doc": normalize_str(row.get("dni", "")), "src": "INVENTARIO_ID"}

    # Priority 2: reparticion_vaquita
    if uname and uname in rep_by_username:
        row = rep_by_username[uname]
        return {"name": normalize_str(row.get("nombres", "")), "doc": normalize_str(row.get("documento", "")), "src": "REPARTICION"}

    # Priority 3: usuarios_funlidi
    if uname and uname in bot_by_username:
        row = bot_by_username[uname]
        return {"name": normalize_str(row.get("nombres_completos", "")), "doc": normalize_str(row.get("numero_documento", "")), "src": "BOT"}
    if user.id > 0 and user.id in bot_by_user_id:
        row = bot_by_user_id[user.id]
        return {"name": normalize_str(row.get("nombres_completos", "")), "doc": normalize_str(row.get("numero_documento", "")), "src": "BOT_ID"}

    # Priority 4: ayudas_humanitarias
    if uname and uname in ayu_by_username:
        row = ayu_by_username[uname]
        return {"name": normalize_str(row.get("nombre", "")), "doc": normalize_str(row.get("dni", "")), "src": "AYUDAS"}

    # Priority 5: Farley CRM
    if uname and uname in farley_by_telegram:
        crm = farley_by_telegram[uname]
        return {"name": normalize_str(crm.get("name", "")), "doc": normalize_str(crm.get("cedula", "")), "src": "FARLEY"}

    # Priority 6: newlist fallback
    if uname and uname in newlist_by_telegram:
        nl = newlist_by_telegram[uname]
        return {"name": normalize_str(nl.get("nombre", "")), "doc": normalize_str(nl.get("documento", "")), "src": "NEWLIST"}

    return None

# ========== MAIN ==========

def main():
    print("\nConectando con Telegram...")
    app = Client("inventario_session", api_id=API_ID, api_hash=API_HASH)

    users = []
    reported_count = 0
    with app:
        users, reported_count = fetch_all_members(app, GROUP_ID)

    total_fetched = len(users)
    print(f"\n  Total obtenidos: {total_fetched}")
    if reported_count and total_fetched < reported_count:
        print(f"  [!] Telegram reporta {reported_count}, pero solo se obtuvieron {total_fetched}")

    matched = 0
    unmatched = 0
    results = []
    for u in users:
        info = lookup_member(u)
        if info:
            matched += 1
        else:
            unmatched += 1
            info = {"name": "", "doc": ""}

        uname_display = f"@{u.username}" if u.username else f"ID:{u.id}"
        real_name = f"{u.first_name or ''} {u.last_name or ''}".strip()

        results.append({
            "telegram_user": uname_display,
            "user_id": u.id,
            "telegram_name": real_name,
            "full_name": info["name"],
            "documento": info["doc"],
            "source": info.get("src", "—") if info else "—",
        })

    results.sort(key=lambda r: r["telegram_user"].lower())

    sep = "=" * 80
    print(f"\n{sep}")
    print(f"  MATCHED: {matched}  |  UNMATCHED: {unmatched}  |  TOTAL: {len(results)}")
    print(f"{sep}")

    # ========== GENERATE XLSX ==========
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("\nERROR: openpyxl no instalado.  pip install openpyxl")
        sys.exit(1)

    wb = Workbook()
    ws = wb.active
    ws.title = "Miembros Grupo Telegram"

    headers = ["Telegram User", "User ID", "Nombre Telegram", "Nombre Completo", "Identificacion/DNI", "Fuente"]
    ws.append(headers)

    header_fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")
    header_font = Font(bold=True, size=11, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    for r in results:
        ws.append([
            r["telegram_user"],
            r["user_id"],
            r["telegram_name"],
            r["full_name"] if r["full_name"] else "—",
            r["documento"] if r["documento"] else "—",
            r["source"],
        ])

    for column in ws.columns:
        max_len = 0
        col_letter = column[0].column_letter
        for cell in column:
            val = str(cell.value) if cell.value else ""
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

    output_path = str(BASE_DIR / "telegram_group_members.xlsx")
    wb.save(output_path)
    print(f"\n  Archivo generado: {output_path}")
    print(f"  {matched} coincidencias, {unmatched} sin datos")
    print(f"{sep}")


if __name__ == "__main__":
    main()

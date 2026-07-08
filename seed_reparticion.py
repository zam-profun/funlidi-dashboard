import os
import sys
from dotenv import load_dotenv
from supabase import create_client, Client
import openpyxl

load_dotenv()

INVENTARIO_SUPABASE_URL = os.environ.get("INVENTARIO_SUPABASE_URL") or os.environ["SUPABASE_URL"]
INVENTARIO_SUPABASE_KEY = os.environ.get("INVENTARIO_SUPABASE_SERVICE_KEY") or os.environ["SUPABASE_SERVICE_KEY"]
SUPABASE_URL = os.environ["SUPABASE_URL"]
SUPABASE_KEY = os.environ["SUPABASE_SERVICE_KEY"]
AYUDAS_SUPABASE_URL = os.environ.get("AYUDAS_SUPABASE_URL") or os.environ["SUPABASE_URL"]
AYUDAS_SUPABASE_KEY = os.environ.get("AYUDAS_SUPABASE_SERVICE_KEY") or os.environ["SUPABASE_SERVICE_KEY"]

supabase_inventario: Client = create_client(INVENTARIO_SUPABASE_URL, INVENTARIO_SUPABASE_KEY)
supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)
supabase_ayudas: Client = create_client(AYUDAS_SUPABASE_URL, AYUDAS_SUPABASE_KEY)

EXCEL_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data to connect", "REPARTICION VAQUITA 100 CUPOS.xlsx")
EXCEL_PATH = os.path.normpath(EXCEL_PATH)

def clean_telegram(val):
    if not val:
        return None
    val = str(val).strip()
    if val.startswith("@@"):
        val = val[1:]
    if val.startswith("@"):
        return val
    return "@" + val

def lookup_personal_info(telegram_username):
    if not telegram_username:
        return (None, None, None)
    tg_lower = telegram_username.lstrip("@").strip().lower()

    # 1. Check inventario_adquisiciones (same INVENTARIO DB)
    try:
        res = supabase_inventario.table("inventario_adquisiciones").select("*").ilike("telegram_username", tg_lower).execute()
        if res.data and len(res.data) > 0:
            r = res.data[0]
            nombre = r.get("nombre") or None
            dni = r.get("dni") or None
            pais = r.get("pais") or None
            if any([nombre, dni, pais]):
                return (nombre, dni, pais)
    except Exception:
        pass

    # 2. Check usuarios_funlidi (main DB)
    try:
        res = supabase.table("usuarios_funlidi").select("*").ilike("telegram_username", tg_lower).execute()
        if res.data and len(res.data) > 0:
            r = res.data[0]
            nombre = r.get("nombres_completos") or None
            dni = r.get("numero_documento") or None
            if any([nombre, dni]):
                return (nombre, dni, None)
    except Exception:
        pass

    # 3. Check ayudas_humanitarias (AYUDAS DB)
    try:
        res = supabase_ayudas.table("ayudas_humanitarias").select("*").ilike("telegram_username", tg_lower).execute()
        if res.data and len(res.data) > 0:
            r = res.data[0]
            nombre = r.get("nombre") or None
            dni = r.get("dni") or None
            pais = r.get("pais") or None
            if any([nombre, dni, pais]):
                return (nombre, dni, pais)
    except Exception:
        pass

    # 4. Check clientes (AYUDAS DB - from telegram-bot-base)
    try:
        res = supabase_ayudas.table("clientes").select("*").ilike("telegram_username", tg_lower).execute()
        if res.data and len(res.data) > 0:
            r = res.data[0]
            nombre = r.get("nombre_completo") or None
            dni = r.get("documento") or None
            if any([nombre, dni]):
                return (nombre, dni, None)
    except Exception:
        pass

    return (None, None, None)

def main():
    if not os.path.exists(EXCEL_PATH):
        print(f"ERROR: Excel file not found at: {EXCEL_PATH}")
        sys.exit(1)

    print(f"Reading Excel file: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows_to_insert = []
    matched_count = 0
    unmatched_count = 0

    row_num = 0
    for row in ws.iter_rows(values_only=True):
        row_num += 1
        if row_num == 1:
            continue
        usuario_raw = row[1] if len(row) > 1 else None
        aporte = row[2] if len(row) > 2 else None
        cant_zim = row[3] if len(row) > 3 else None
        cant_dinar = row[4] if len(row) > 4 else None
        cant_oro = row[5] if len(row) > 5 else None
        cajas_total = row[6] if len(row) > 6 else None

        if not usuario_raw or not aporte:
            continue

        telegram_username = clean_telegram(usuario_raw)
        if not telegram_username:
            continue

        nombres, documento, pais = lookup_personal_info(telegram_username)

        status = "MATCHED" if any([nombres, documento, pais]) else "UNMATCHED"
        if status == "MATCHED":
            matched_count += 1
        else:
            unmatched_count += 1

        print(f"  [{status}] {telegram_username} | Aporte: {aporte} | ZIM: {cant_zim} | DINAR: {cant_dinar} | ORO: {cant_oro} | Total: {cajas_total} | Nombre: {nombres or '-'} | Doc: {documento or '-'} | Pais: {pais or '-'}")

        rows_to_insert.append({
            "telegram_username": telegram_username,
            "aporte": int(aporte) if aporte else 0,
            "cant_zim": int(cant_zim) if cant_zim else 0,
            "cant_dinar": int(cant_dinar) if cant_dinar else 0,
            "cant_oro": int(cant_oro) if cant_oro else 0,
            "cajas_total": int(cajas_total) if cajas_total else 0,
            "nombres": nombres,
            "documento": documento,
            "pais": pais,
        })

    wb.close()

    print(f"\nTotal rows to insert: {len(rows_to_insert)}")
    print(f"Matched: {matched_count}")
    print(f"Unmatched: {unmatched_count}")

    if rows_to_insert:
        print("\nInserting into Supabase table 'reparticion_vaquita'...")
        result = supabase_inventario.table("reparticion_vaquita").insert(rows_to_insert).execute()
        print(f"Inserted {len(result.data)} rows successfully!")

        if len(result.data) != len(rows_to_insert):
            print(f"WARNING: Expected {len(rows_to_insert)} but got {len(result.data)}")

    print("\nDone!")

if __name__ == "__main__":
    main()
